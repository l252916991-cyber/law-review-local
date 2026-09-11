import io
import os
import tempfile
import unittest
from pathlib import Path

from tests.support import IsolatedDatabaseTestCase


class BankTransactionParserTest(unittest.TestCase):
    def test_csv_aliases_decimal_direction_and_utf8_bom(self):
        from app.bank_transactions import parse_csv
        payload = "\ufeff账号,收支方向,交易金额,交易日期,对方户名,摘要\nA,收入,\"1,234.50\",2024-01-02,B,款项".encode()
        rows = parse_csv(payload)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].amount_minor, 123450)
        self.assertEqual(rows[0].direction, "inflow")
        self.assertEqual(rows[0].transaction_time, "2024-01-02")
        self.assertEqual(rows[0].parse_status, "parsed")

    def test_gb18030_and_needs_review_warnings(self):
        from app.bank_transactions import parse_csv
        payload = "账户,金额,日期,对方账户\n甲,坏金额,日期,乙\n".encode("gb18030")
        rows = parse_csv(payload)
        self.assertEqual(rows[0].parse_status, "needs_review")
        self.assertIn("invalid_amount", rows[0].warnings)
        self.assertIn("invalid_date", rows[0].warnings)
        self.assertIn("missing_direction", rows[0].warnings)

    def test_limits_rows_and_cells(self):
        from app.bank_transactions import parse_csv
        with self.assertRaises(ValueError):
            parse_csv(("账号,金额\n" + "a,1\n" * 100001).encode())

    def test_xlsx_uses_value_only_parser_and_keeps_sheet_provenance(self):
        from datetime import date

        from openpyxl import Workbook

        from app.bank_transactions import parse_spreadsheet

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "流水"
        sheet.append(["账号", "收支方向", "交易金额", "交易日期", "对方户名", "摘要"])
        sheet.append(["A", "收入", 12.5, date(2024, 1, 2), "B", "入账"])
        formula_sheet = workbook.create_sheet("公式")
        formula_sheet.append(["账号", "收支方向", "交易金额", "交易日期", "对方户名"])
        formula_sheet.append(["A", "收入", "=1+1", date(2024, 1, 3), "C"])
        payload = io.BytesIO()
        workbook.save(payload)

        rows = parse_spreadsheet(payload.getvalue(), "流水.xlsx")
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0].source_sheet, "流水")
        self.assertEqual(rows[0].amount_minor, 1250)
        self.assertEqual(rows[0].transaction_time, "2024-01-02")
        # Formula cells are not executed; without a cached value they remain a
        # review item instead of being treated as a calculated amount.
        self.assertIsNone(rows[1].amount_minor)
        self.assertIn("missing_amount_value", rows[1].warnings)

    def test_spreadsheet_payload_limit_is_checked_before_parsing(self):
        from app.bank_transactions import MAX_PAYLOAD_BYTES, parse_spreadsheet

        with self.assertRaises(ValueError):
            parse_spreadsheet(b"x" * (MAX_PAYLOAD_BYTES + 1), "流水.xlsx")

    def test_malformed_spreadsheet_is_a_bounded_value_error(self):
        from app.bank_transactions import parse_spreadsheet

        with self.assertRaisesRegex(ValueError, "无法解析"):
            parse_spreadsheet(b"not-a-workbook", "流水.xlsx")


class BankTransactionServiceTest(IsolatedDatabaseTestCase):
    @classmethod
    def setUpClass(cls):
        from app.db import init_db
        init_db(seed=True)

    def test_persist_summary_graph_and_replay_are_exact(self):
        from app.db import connect
        from app.services import bank_transaction_graph, list_bank_transactions, persist_bank_transactions, summarize_bank_transactions
        payload = "账号,收支方向,交易金额,交易日期,对方户名,摘要\nA,收入,100.00,2024-01-02,B,入账\nA,支出,20.00,2024-01-03,C,付款\n".encode()
        with tempfile.TemporaryDirectory(prefix="bank-source-") as directory:
            source = Path(directory) / "流水.csv"
            source.write_bytes(payload)
            with connect() as conn:
                document_id = conn.execute("SELECT id FROM documents WHERE case_id=1 LIMIT 1").fetchone()[0]
                source_hash = conn.execute("SELECT content_hash FROM documents WHERE id=?", (document_id,)).fetchone()[0] or "source-hash"
            first = persist_bank_transactions(1, document_id, source_hash, payload)
            second = persist_bank_transactions(1, document_id, source_hash, payload)
        self.assertEqual(first["rows"], 2)
        self.assertEqual(second["rows"], 2)
        rows = list_bank_transactions(1)["transactions"]
        self.assertEqual(len(rows), 2)
        summary = summarize_bank_transactions(1)
        self.assertEqual(summary["inflow_minor"], 10000)
        self.assertEqual(summary["outflow_minor"], 2000)
        self.assertEqual(summary["net_minor"], 8000)
        graph = bank_transaction_graph(1)
        self.assertEqual(graph["transaction_count"], 2)
        self.assertTrue(all(edge["relation_type"] == "资金链路" for edge in graph["edges"]))
        self.assertTrue(any(edge["from"] == "party:B" and edge["to"] == "account:A" for edge in graph["edges"]))
        with connect() as conn:
            self.assertEqual(conn.execute("SELECT COUNT(*) FROM bank_transactions WHERE case_id=1").fetchone()[0], 2)

    def test_case_delete_cascades_transactions(self):
        from app.db import connect, now, transaction
        with transaction() as conn:
            case_id = conn.execute("INSERT INTO cases(title,created_at,updated_at) VALUES ('流水案件',?,?)", (now(), now())).lastrowid
            conn.execute("INSERT INTO bank_transactions(case_id,source_row_number,account,parser_version,row_fingerprint,created_at) VALUES (?,?,?,?,?,?)", (case_id, 2, "A", "test", "fingerprint", now()))
        with transaction() as conn:
            conn.execute("DELETE FROM cases WHERE id=?", (case_id,))
        with connect() as conn:
            self.assertEqual(conn.execute("SELECT COUNT(*) FROM bank_transactions WHERE case_id=?", (case_id,)).fetchone()[0], 0)

    def test_xlsx_upload_is_indexed_and_persisted_with_sheet_reference(self):
        from datetime import date

        from openpyxl import Workbook

        from app.db import connect
        from app.services import index_upload, parse_spreadsheet, persist_parsed_bank_rows

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "主表"
        sheet.append(["账号", "收支方向", "交易金额", "交易日期", "对方户名"])
        sheet.append(["A", "支出", 20, date(2024, 2, 1), "C"])
        payload = io.BytesIO()
        workbook.save(payload)
        content = payload.getvalue()

        document = index_upload(1, "流水.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        rows = parse_spreadsheet(content, "流水.xlsx")
        result = persist_parsed_bank_rows(1, document["id"], document["content_hash"], rows)

        self.assertEqual(document["pages"], 1)
        self.assertEqual(result["parsed"], 1)
        with connect() as conn:
            stored = conn.execute("SELECT source_sheet, amount_minor, direction FROM bank_transactions WHERE source_document_id=?", (document["id"],)).fetchone()
        self.assertEqual(tuple(stored), ("主表", 2000, "outflow"))
